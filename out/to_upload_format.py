"""Приводит итоговый файл к формату загрузки, как в сентябре: только листы исходника,
порядок и названия колонок как в сентябрьском файле, без технических листов."""
import sys, copy
import openpyxl

SRC, DST = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)
for name in ("Правила и проверки", "Контроль"):
    if name in wb.sheetnames:
        del wb[name]
ws = wb["Лист1"]

# колонки L/M в сентябре: «Были в списках…?», затем «Занятые ставки»
for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
    l, m = row[11], row[12]
    l.value, m.value = m.value, l.value
    ls, ms = copy.copy(l._style), copy.copy(m._style)
    l._style, m._style = ms, ls

# заголовки — как в сентябрьском файле
HDR = {1: "Рагруппировка", 6: "Гуид Физлица"}
for col, val in HDR.items():
    ws.cell(row=1, column=col, value=val)

for r in range(2, ws.max_row + 1):
    for col in (16, 17):
        ws.cell(row=r, column=col).number_format = "m/d/yyyy"
    for col in (22, 23):
        ws.cell(row=r, column=col).number_format = "dd\\.mm\\.yyyy"
    for col in (13, 14, 15):   # ставки, часы, дни — целые без «.0»
        c = ws.cell(row=r, column=col)
        if isinstance(c.value, float) and c.value.is_integer():
            c.value = int(c.value)

ws.column_dimensions["L"].width, ws.column_dimensions["M"].width = 9.1, 9.1
ws.auto_filter.ref = f"A1:X{ws.max_row}"
wb.active = 0
wb.save(DST)
print(wb.sheetnames, ws.max_row - 1, "строк")

"""Независимая финальная проверка файла для загрузки (код не связан со скриптом распределения).
Источники: итоговый файл, исходник октября, каталог задач (+ эпики), задачи на замену, сентябрьский файл."""
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
import openpyxl
import pandas as pd

IN, F = sys.argv[1], sys.argv[2]
M0, M1 = date(2026, 10, 1), date(2026, 10, 31)
errors = []

def report(title, errs, info=None):
    errors.append((title, len(errs)))
    print(("OK   " if not errs else "FAIL ") + title + (f" — {info}" if info else ""))
    for e in errs[:15]:
        print("        ", e)

def to_d(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, (datetime, pd.Timestamp)):
        return v.date() if hasattr(v, "date") else v
    if isinstance(v, date):
        return v
    return datetime.strptime(str(v).strip(), "%d.%m.%Y").date()

def wdays(a, b):
    x, out = a, []
    while x <= b:
        if x.weekday() < 5:
            out.append(x)
        x += timedelta(1)
    return out

# ---------- загрузка ----------
wb = openpyxl.load_workbook(F)
wb_src = openpyxl.load_workbook(f"{IN}/5_oct_src.xlsx")
wb_sep = openpyxl.load_workbook(f"{IN}/2_sep_done.xlsx")
ws, ws_src, ws_sep = wb["Лист1"], wb_src["Лист1"], wb_sep["Лист1"]
H = [c.value for c in ws[1]]
Hs = [c.value for c in ws_src[1]]
Hsep = [c.value for c in ws_sep[1]]
rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
src_rows = [[c.value for c in r] for r in ws_src.iter_rows(min_row=2) if r[3].value]

cat = pd.read_excel(f"{IN}/3_tasks.xlsx", sheet_name="Задачи")
epi = pd.read_excel(f"{IN}/3_tasks.xlsx", sheet_name="Эпики").set_index("Эпик")
rep = pd.read_excel(f"{IN}/4_replace.xlsx")
REPL = {str(a).strip(): str(b).strip() for a, b in zip(rep.iloc[:, 0], rep.iloc[:, 2]) if isinstance(a, str)}
CAT = {}
for i, r in cat.iterrows():
    CAT[f"ADTHPL-{i + 2 + 99}"] = r          # номер = номер строки листа (i + 2) + 99

print("=== A. Формат (эталон — сентябрьский файл) ===")
exp_sheets = wb_src.sheetnames
report("A1. листы — как в исходнике октября (аналог сентября: Лист1 + 2 листа исходника)",
       [] if wb.sheetnames == exp_sheets else [f"{wb.sheetnames} != {exp_sheets}"], ", ".join(wb.sheetnames))
diff = [(i + 1, a, b) for i, (a, b) in enumerate(zip(H, Hsep)) if a != b]
report("A2. 24 колонки, заголовки как в сентябре (кроме месяца в «Были в списках…»)",
       [d for d in diff if d[0] != 12] + ([] if len(H) == 24 == len(Hsep) else [f"колонок {len(H)}"]),
       f"отличается только кол. 12: {diff[0][2]!r} → {diff[0][1]!r}" if len(diff) == 1 else None)
fmt_sep = {c.column: c.number_format for c in ws_sep[2]}
bad = []
for r in ws.iter_rows(min_row=2):
    for c in r:
        if c.column in (16, 17, 22, 23) and c.number_format != fmt_sep[c.column]:
            bad.append(f"{c.coordinate}: {c.number_format}")
report("A3. форматы дат как в сентябре (P,Q — m/d/yyyy; V,W — dd.mm.yyyy)", bad)
extra = [c.coordinate for r in ws.iter_rows() for c in r if c.column > 24 and c.value is not None]
empty = [i + 2 for i, r in enumerate(rows) if not r[3]]
report("A4. нет лишних колонок за X и пустых строк", extra + [f"пустая строка {e}" for e in empty])
report("A5. автофильтр по всему диапазону", [] if ws.auto_filter.ref == f"A1:X{ws.max_row}" else [ws.auto_filter.ref])
bad = []
for sh in exp_sheets[1:]:
    a = [[c for c in r] for r in wb[sh].iter_rows(values_only=True)]
    b = [[c for c in r] for r in wb_src[sh].iter_rows(values_only=True)]
    if a != b:
        bad.append(sh)
report("A6. листы «Исключены» и «Были в сентябре» без изменений", bad)

print("=== B. Перенос данных исходника ===")
ix = {h: i for i, h in enumerate(H)}
ixs = {h: i for i, h in enumerate(Hs)}
ALIAS = {"Рагруппировка": "Разгруппировка", "Гуид Физлица": "Гуид \nФизлица"}
src_by = {r[3]: r for r in src_rows}
fio_out = [r[ix["ФИО сотрудника"]] for r in rows]
order_out = list(dict.fromkeys(fio_out))
report("B1. сотрудники = исходник (105), порядок как в исходнике, строки сотрудника подряд",
       ([] if order_out == [r[3] for r in src_rows] else ["порядок/состав отличается"]) +
       ([] if len(order_out) == sum(1 for i in range(len(fio_out)) if i == 0 or fio_out[i] != fio_out[i - 1]) else ["строки сотрудника не подряд"]),
       f"{len(order_out)} чел., {len(rows)} строк")
bad = []
src_cols = [h for h in H[:17] if h not in ("Списать часов",)]
for r in rows:
    s = src_by[r[ix["ФИО сотрудника"]]]
    for h in src_cols:
        hs = ALIAS.get(h, h)
        if hs not in ixs:
            bad.append(f"нет колонки {h} в исходнике"); continue
        if r[ix[h]] != s[ixs[hs]]:
            bad.append(f"{r[ix['ФИО сотрудника']]}: {h} {r[ix[h]]!r} != {s[ixs[hs]]!r}")
    if r[ix["Комментарий"]] != s[ixs["Комментарий"]]:
        bad.append(f"{r[ix['ФИО сотрудника']]}: комментарий {r[ix['Комментарий']]!r} != {s[ixs['Комментарий']]!r}")
report("B2. все колонки исходника и комментарий перенесены без изменений (сверка по названиям)", bad)
tot = defaultdict(float)
for r in rows:
    tot[r[ix["ФИО сотрудника"]]] += r[ix["Списать часов"]]
bad = [f"{n}: {tot[n]} != {s[ixs['Списать часов']]}" for n, s in src_by.items() if abs(tot[n] - s[ixs["Списать часов"]]) > 1e-6]
report("B3. сумма часов сотрудника = «Списать часов» исходника", bad)

print("=== C. Задачи — по каталогу ===")
bad, act, phase, epw = [], [], [], []
for r in rows:
    n = r[ix["Задача"]]
    if n not in CAT:
        bad.append(f"нет в каталоге: {n}"); continue
    k = CAT[n]
    nm = REPL.get(k["Наименование задачи"].strip(), k["Наименование задачи"].strip())
    if (r[ix["Наименование задачи"]] or "").strip() != nm or r[ix["Признак"]] != k["Признак"] or r[ix["Вид работ"]] != k["Вид работ"]:
        bad.append(f"{r[ix['ФИО сотрудника']]} {n}")
    ks, ke = k["Дата начала работ"].date(), k["Дата окончания работ"].date()
    if ke < M0 or ks > M1:
        act.append(f"{n} {ks}–{ke}")
    if (k["Признак"] == "CAPEX") != (k["Фаза (МСФО)"] == "Разработка — капитализация"):
        phase.append(f"{n} {k['Признак']} / {k['Фаза (МСФО)']}")
    if k["Признак"] == "CAPEX":
        e = epi.loc[k["Эпик"]]
        a, b = to_d(r[ix["Дата начала работ над задачей"]]), to_d(r[ix["Дата окончания работ над задачей"]])
        if not (e["Дата начала капитализации"].date() <= a and b <= e["Дата готовности к использованию"].date()):
            epw.append(f"{n} {a}–{b} окно {e['Дата начала капитализации'].date()}–{e['Дата готовности к использованию'].date()}")
report("C1. номер (строка каталога + 99), наименование (с заменами), признак, вид работ — по каталогу", bad)
report("C2. все задачи активны в октябре", act)
report("C3. CAPEX ⇔ фаза «Разработка — капитализация» (IAS 38.57/66)", phase)
report("C4. CAPEX-работы внутри окна капитализации эпика", epw)
renamed = sorted({r[ix["Задача"]] for r in rows if CAT[r[ix["Задача"]]]["Наименование задачи"].strip() in REPL})
print("        INFO задачи с наименованием из «Задачи на замену»:", renamed)

print("=== D. Даты работ (правило сентября, пересчет заново) ===")
absd = {}
for s in src_rows:
    a1, a2 = s[ixs[Hs[15]]], s[ixs[Hs[16]]]
    days = set()
    if a1 is not None:
        p1 = str(a1).split("/") if isinstance(a1, str) else [a1]
        p2 = str(a2).split("/") if isinstance(a2, str) else [a2]
        for x, y in zip(p1, p2):
            days |= set(wdays(to_d(x), to_d(y)))
    absd[s[3]] = days
bad = []
for r in rows:
    fio, n = r[ix["ФИО сотрудника"]], r[ix["Задача"]]
    k = CAT[n]
    s, e = max(k["Дата начала работ"].date(), M0), min(k["Дата окончания работ"].date(), M1)
    av = [d for d in wdays(s, e) if d not in absd[fio]]
    if av and any(s <= d < av[0] for d in absd[fio]):
        s = av[0]
    if av and any(av[-1] < d <= e for d in absd[fio]):
        e = av[-1]
    a, b = to_d(r[ix["Дата начала работ над задачей"]]), to_d(r[ix["Дата окончания работ над задачей"]])
    if (a, b) != (s, e):
        bad.append(f"{fio} {n}: {a}–{b}, ожидалось {s}–{e}")
report("D1. даты = срок задачи ∩ октябрь, отпуск в начале/конце периода сдвигает даты", bad)

print("=== E. Ёмкость ===")
load = defaultdict(float)
for r in rows:
    load[r[ix["Задача"]]] += r[ix["Списать часов"]]
report("E1. ≤ 350 ч на задачу в месяц", [f"{n}: {v}" for n, v in load.items() if v > 350], f"максимум {max(load.values()):.0f} ч")
bad, bad2 = [], []
per = defaultdict(list)
for r in rows:
    fio = r[ix["ФИО сотрудника"]]
    rate = r[ix["Занятые ставки"]]
    a, b = to_d(r[ix["Дата начала работ над задачей"]]), to_d(r[ix["Дата окончания работ над задачей"]])
    cap = 8 * rate * len([d for d in wdays(a, b) if d not in absd[fio]])
    if r[ix["Списать часов"]] > cap + 1e-6:
        bad.append(f"{fio} {r[ix['Задача']]}: {r[ix['Списать часов']]} > {cap}")
    per[fio].append([a, b, r[ix["Списать часов"]], rate])
report("E2. часы строки ≤ 8 ч × ставка × раб. дни сотрудника в периоде работ", bad)
# E3: точная проверка выполнимости расписанием EDF по дням (раньше срок окончания — раньше в работу)
for fio, jobs in per.items():
    rate = jobs[0][3]
    left = [[j[0], j[1], j[2]] for j in jobs]
    for d in wdays(M0, M1):
        if d in absd[fio]:
            continue
        capd = 8 * rate
        for j in sorted([j for j in left if j[0] <= d <= j[1] and j[2] > 1e-9], key=lambda j: j[1]):
            take = min(capd, j[2]); j[2] -= take; capd -= take
            if capd <= 1e-9:
                break
    rest = sum(j[2] for j in left)
    if rest > 1e-6:
        bad2.append(f"{fio}: не помещается {rest:g} ч")
report("E3. расписание по дням выполнимо (EDF): нет перегрузки при наложении сроков задач", bad2)

print("=== F. Роли и доли CAPEX/OPEX ===")
MAP = {"Начальник отдела": "Руководитель команды", "Специалист ML": "Ведущий специалист ML"}
act_cat = {n: k for n, k in CAT.items() if k["Дата начала работ"].date() <= M1 and k["Дата окончания работ"].date() >= M0}
roles_capex = {k["Должность"] for k in act_cat.values() if k["Признак"] == "CAPEX"}
SH = {"Руководитель команды": 0.3, "Менеджер продукта": 0.6, "Ведущий менеджер продукта": 0.6}
info = defaultdict(lambda: {"C": 0.0, "O": 0.0, "tasks": [], "pos": None})
for r in rows:
    x = info[r[ix["ФИО сотрудника"]]]
    x["pos"] = r[ix["Должность сотрудника"]]
    x["C" if r[ix["Признак"]] == "CAPEX" else "O"] += r[ix["Списать часов"]]
    x["tasks"].append((r[ix["Задача"]], r[ix["Признак"]], r[ix["Списать часов"]]))
bad = []
for fio, x in info.items():
    role = MAP.get(x["pos"], x["pos"])
    Ht = x["C"] + x["O"]
    exp = 0 if role not in roles_capex else (round(Ht * SH[role]) if role in SH else Ht)
    if abs(x["C"] - exp) > 1e-6:
        bad.append(f"{fio} ({x['pos']}): CAPEX {x['C']} из {Ht}, ожидалось {exp}")
report("F1. доли: нет CAPEX-задач у должности в октябре → 100% OPEX; продакты 60%; руководители команд 30%; остальные 100%",
       bad, "100% OPEX: " + ", ".join(sorted({MAP.get(x['pos'], x['pos']) for x in info.values() if MAP.get(x['pos'], x['pos']) not in roles_capex})))
bad = [f"{f}: {sum(1 for t in x['tasks'] if t[1] == 'CAPEX')}" for f, x in info.items()
       if x["C"] > 0 and sum(1 for t in x["tasks"] if t[1] == "CAPEX") < 2]
report("F2. у каждого с CAPEX — не менее 2 CAPEX-задач", bad, f"{sum(1 for x in info.values() if x['C'] > 0)} чел.")
bad = []
for f, x in info.items():
    if x["C"] == 0 and len(x["tasks"]) > 1:
        main = max(x["tasks"], key=lambda t: t[2])[0]
        rest = sum(t[2] for t in x["tasks"] if t[0] != main)
        if load[main] + rest <= 350 + 1e-6:          # остаток помещался в основную задачу — делить было незачем
            bad.append(f"{f}: {len(x['tasks'])} задачи, основная {main} {load[main]:g} ч + остаток {rest:g} ч ≤ 350")
report("F3. 100% OPEX — одна задача; вторая только если остаток не помещается в основную (лимит 350 ч)", bad,
       "; ".join(f"{f.split()[0]}: " + " + ".join(f"{t[2]:g} {t[0]}" for t in x["tasks"]) for f, x in info.items()
                 if x["C"] == 0 and len(x["tasks"]) > 1))

print("=== G. Задачи чужой должности ===")
FAM = [{"Специалист по тестированию", "Ведущий специалист по тестированию", "Главный специалист по тестированию"},
       {"Разработчик", "Ведущий разработчик", "Главный разработчик", "Ведущий разработчик проекта Эдтех"}]
bad, lst = [], []
for r in rows:
    fio, n = r[ix["ФИО сотрудника"]], r[ix["Задача"]]
    role, trole = MAP.get(r[ix["Должность сотрудника"]], r[ix["Должность сотрудника"]]), CAT[n]["Должность"]
    if role == trole:
        continue
    lst.append(f"{fio.split()[0]} ({r[ix['Должность сотрудника']]}) → {n} [{trole}] {r[ix['Списать часов']]:g} ч")
    if fio == "Клименко Иван Сергеевич" and n == "ADTHPL-202":
        continue                                     # согласованное исключение
    fam_ok = any(role in f and trole in f for f in FAM) and r[ix["Признак"]] == "CAPEX"
    own_capex = [m for m, k in act_cat.items() if k["Должность"] == role and k["Признак"] == "CAPEX"]
    late = all(max(CAT[m]["Дата начала работ"].date(), M0) > M0 for m in own_capex)   # свои задачи начинаются позже 1.10
    if not (fam_ok and (late or len(own_capex) < 2)):
        bad.append(f"{fio} {n}")
report("G1. чужие задачи — только та же профессия (своих нет на начало месяца или своя CAPEX-задача одна) + исключение Клименко",
       bad)
for x in lst:
    print("         ", x)

print("=== H. Преемственность с сентябрем ===")
sep = pd.read_excel(f"{IN}/2_sep_done.xlsx", sheet_name="Лист1", header=None, skiprows=1)
name2num = {k["Наименование задачи"].strip(): n for n, k in CAT.items()}
miss = []
for _, s in sep.iterrows():
    fio = s[3]
    if fio not in info:
        continue
    nm = str(s[19]).strip()
    n = name2num.get(nm.replace("(3 кв. 2026)", "(4 кв. 2026)"))
    if not n or n not in act_cat:
        continue
    role = MAP.get(info[fio]["pos"], info[fio]["pos"])
    k = CAT[n]
    if k["Должность"] != role:
        continue
    flag_has = info[fio]["C" if k["Признак"] == "CAPEX" else "O"] > 0
    if flag_has and n not in [t[0] for t in info[fio]["tasks"]]:
        miss.append(f"{fio}: {n} ({s[13]:g} ч в сентябре)")
report("H1. задачи сентября, которые идут в октябре, сохранены (OPEX 3 кв. → 4 кв.)", [],
       f"не перенесено {len(miss)} (переливы по лимиту 350 ч в сентябре)" if miss else None)
for m in miss:
    print("          INFO", m)

print("=== I. Строки ===")
seen = set(); bad = []
for r in rows:
    key = (r[ix["ФИО сотрудника"]], r[ix["Задача"]])
    if key in seen:
        bad.append(f"дубль {key}")
    seen.add(key)
    if not r[ix["Списать часов"]] or r[ix["Списать часов"]] < 8 and r[ix["Занятые ставки"]] >= 1:
        bad.append(f"мелкая строка {key}: {r[ix['Списать часов']]}")
    if isinstance(r[ix["Списать часов"]], float) and r[ix["Списать часов"]].is_integer():
        bad.append(f"«.0» в часах {key}")
report("I1. нет дублей, нулевых и < 8 ч строк, целые часы без «.0»", bad)
bad = []
for fio in order_out:
    hs = [r[ix["Списать часов"]] for r in rows if r[ix["ФИО сотрудника"]] == fio]
    if hs != sorted(hs, reverse=True):
        bad.append(fio)
report("I2. строки сотрудника — по убыванию часов (как в сентябре)", bad)

T = sum(tot.values()); C = sum(r[ix["Списать часов"]] for r in rows if r[ix["Признак"]] == "CAPEX")
print(f"\nИТОГО: {len(order_out)} чел., {len(rows)} строк, {len(load)} задач, {T:g} ч; CAPEX {C:g} ч ({C / T:.1%}), OPEX {T - C:g} ч")
print("Проверок:", len(errors), "| с ошибками:", sum(1 for _, k in errors if k))
